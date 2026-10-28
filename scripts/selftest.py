"""自测脚本：不依赖 LM Studio，验证解析、检索、抽取校验、Excel 写回全链路。

用法：
    python scripts/selftest.py
"""

from __future__ import annotations

import copy
import json
import os
import shutil
import sys
import tempfile
import time

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pymupdf as fitz
import openpyxl

from rag import config as cfg
from rag.extract.extractor import Extractor
from rag.extract.normalize import canonical_unit, parse_value, verify_quote
from rag.extract.schema import build_schema_from_headers
from rag.index import Embedder, Store
from rag.index.retriever import Retriever
from rag.ingest import parse_pdf
from rag.models import Chunk, Evidence, FieldValue, MaterialRecord
from rag.writeback import diff as diffmod
from rag.writeback.excel import ExcelTarget

PASS, FAIL = [], []


def check(name: str, cond: bool, detail: str = "") -> None:
    (PASS if cond else FAIL).append(name)
    print(f"  [{'PASS' if cond else 'FAIL'}] {name}" + (f" | {detail}" if detail else ""))


# ---------------------------------------------------------------- 1. 数值解析

def t_parse_value():
    print("\n[1] 数值与区间解析（模型不做算术，全部由此处确定性代码完成）")
    cases = [
        ("0.55", "single", 0.55),
        ("0.5~0.6", "range", None),
        ("≤0.05", "upper_bound", None),
        ("≥300", "lower_bound", None),
        # 半角比较符必须和 ≤/≥ 同级处理：漏了它们会把 "<0.05" 当成实测值 0.05
        ("<0.05", "upper_bound", None),
        (">300", "lower_bound", None),
        ("<=0.05", "upper_bound", None),
        (">=300", "lower_bound", None),
        ("余量", "text", None),
    ]
    for raw, kind, val in cases:
        r = parse_value(raw)
        ok = r["kind"] == kind and (val is None or r["value"] == val)
        check(f"parse_value({raw!r}) -> {kind}", ok, f"实得 {r}")
    r = parse_value("0.20~0.50")
    check("区间不取中值（low/high 原样保留）", r["low"] == 0.20 and r["high"] == 0.50 and r["value"] is None,
          f"low={r['low']} high={r['high']} value={r['value']}")
    r = parse_value("380 ± 2")
    check("均值±标准差 不当作区间（mean_pm）",
          r["kind"] == "mean_pm" and r["mean"] == 380.0 and r["pm"] == 2.0
          and r["low"] is None and r["high"] is None,
          f"实得 {r}")
    r = parse_value("8.52 ± 0.91")
    check("均值±标准差 小数保留", r["kind"] == "mean_pm" and r["mean"] == 8.52 and r["pm"] == 0.91,
          f"实得 {r}")
    check("单位归一 wt% -> wt.%", canonical_unit("wt%") == "wt.%")
    check("单位归一 ℃ -> °C", canonical_unit("℃") == "°C")


# ---------------------------------------------------------------- 2. 引文校验

def t_verify():
    print("\n[2] 原文证据回链（防止模型编造）")
    text = "The alloy was solution treated at 530 °C for 2 h, then aged at 175 °C for 8 h."
    check("命中真实引文", verify_quote("aged at 175 °C for 8 h", [text]) == 0)
    check("忽略空白差异", verify_quote("aged at 175°C for 8h", [text]) == 0)
    check("拒绝编造引文", verify_quote("aged at 190 °C for 10 h", [text]) == -1)


# ---------------------------------------------------------------- 3. 表头映射

def t_schema():
    print("\n[3] Excel 表头 -> 抽取字段映射")
    heads = ["合金牌号", "Si(wt.%)", "Mg", "固溶温度", "时效温度", "抗拉强度Rm(MPa)", "自定义列XYZ"]
    sch = build_schema_from_headers(heads)
    matched = {m["header"]: m["key"] for m in sch["matched"]}
    check("识别 Si(wt.%) -> composition.Si", matched.get("Si(wt.%)") == "composition.Si")
    check("识别 时效温度 -> process.aging_temp", matched.get("时效温度") == "process.aging_temp")
    check("识别 抗拉强度Rm(MPa) -> properties.UTS", matched.get("抗拉强度Rm(MPa)") == "properties.UTS")
    check("未知列按自定义列保留", any(c["header"] == "自定义列XYZ" for c in sch["custom"]))
    check("单位可从表头推断", sch["specs"]["properties.UTS"].unit == "MPa")


# ---------------------------------------------------------------- 3.5 双栏阅读顺序

def t_two_column(tmp: str):
    print("\n[3.5] 双栏论文阅读顺序（回归测试：左栏下段不得串到右栏上段之前）")
    pdf = os.path.join(tmp, "two_col.pdf")
    doc = fitz.open()
    page = doc.new_page()                      # 约 595x842
    # 左栏两段、右栏两段，左右栏 y 坐标交错重叠，专测行带法的陷阱
    page.insert_text((72, 100), "LEFT AAA first paragraph of left column.", fontsize=10)
    page.insert_text((72, 300), "LEFT BBB second paragraph of left column.", fontsize=10)
    page.insert_text((320, 100), "RIGHT CCC first paragraph of right column.", fontsize=10)
    page.insert_text((320, 300), "RIGHT DDD second paragraph of right column.", fontsize=10)
    doc.save(pdf)
    doc.close()

    chunks, _ = parse_pdf(pdf, cfg.load())
    joined = "\n".join(c.text for c in chunks if c.kind == "text")
    ia, ib, ic = (joined.find(k) for k in ("LEFT AAA", "LEFT BBB", "RIGHT CCC"))
    check("解析出双栏文本", all(x >= 0 for x in (ia, ib, ic)), f"{len(chunks)} 块")
    check("阅读顺序：左栏两段先于右栏", 0 <= ia < ib < ic,
          f"AAA@{ia} BBB@{ib} CCC@{ic}")




def make_test_pdf(path: str) -> None:
    """生成一篇含成分表与工艺段的迷你论文，用于解析与检索测试。"""
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 70), "Composition and ageing behaviour of Al-Mg-Si alloy",
                     fontsize=13)
    page.insert_text((72, 92), "1. Materials and methods", fontsize=10)

    # 表格网格
    x0, y0, rowh = 72, 120, 22
    colw = [110, 70, 70, 70, 70]
    data = [
        ["Alloy", "Si", "Mg", "Cu", "Fe"],
        ["6061-T6", "0.62", "1.02", "0.24", "0.35"],
        ["6063-T5", "0.45", "0.55", "0.02", "0.22"],
    ]
    for r, row in enumerate(data):
        x = x0
        for c, cell in enumerate(row):
            page.draw_rect(fitz.Rect(x, y0 + r * rowh, x + colw[c], y0 + (r + 1) * rowh),
                           color=(0, 0, 0), width=0.6)
            page.insert_text((x + 4, y0 + r * rowh + 15), cell, fontsize=9)
            x += colw[c]

    page.insert_textbox(fitz.Rect(72, 200, 520, 300),
                        "The 6061 alloy was homogenized at 560 C for 6 h. Solution treatment "
                        "was carried out at 530 C for 2 h followed by water quenching. "
                        "Artificial ageing was performed at 175 C for 8 h. "
                        "The yield strength reached 276 MPa and the tensile strength reached "
                        "310 MPa with an elongation of 12 %.",
                        fontsize=10)
    doc.save(path)
    doc.close()


# ---------------------------------------------------------------- 5. 解析 + 索引 + 检索

def t_pipeline(tmp: str):
    print("\n[5] PDF 解析 / 切块 / 索引 / 混合检索")
    pdf = os.path.join(tmp, "test_paper.pdf")
    make_test_pdf(pdf)
    chunks, meta = parse_pdf(pdf, cfg.load())
    n_tab = sum(1 for c in chunks if c.kind == "table")
    check("解析出文本块", any(c.kind == "text" for c in chunks), f"共 {len(chunks)} 块")
    check("识别出表格块", n_tab >= 1, f"表格 {n_tab} 个")
    check("页码信息完整", all(c.page >= 1 for c in chunks))

    idx_dir = os.path.join(tmp, "index")
    store = Store(idx_dir)
    emb = Embedder(cfg.load())
    emb.load()
    # embedding 可能走本机模型（需 torch）或 LM Studio（需服务在线）。
    # 不可用时降级为占位向量，继续验证索引结构，避免整套自测因外部依赖挂掉。
    emb_live = True
    try:
        vecs = emb.encode_documents([c.text for c in chunks])
    except Exception as e:
        emb_live = False
        import numpy as np          # 仅在降级路径需要
        vecs = np.zeros((len(chunks), 8), dtype=np.float32)
        print(f"  [SKIP] embedding 服务不可用（{type(e).__name__}），向量相关断言跳过")
    store.add(chunks, vecs)
    store.set_meta(meta["doc_id"], meta)
    store.save()
    st = store.stats()
    if emb_live:
        # 维度必须与模型实际输出一致（曾写死 512，换 768 维模型后误报）
        check("向量维度与模型一致", st["dim"] == emb.dim and st["dim"] > 0,
              f"索引 dim={st['dim']} 模型 dim={emb.dim}")
    check("向量数量与块数一致", st["vectors"] == st["chunks"])

    if emb_live:
        # 检索也要走编码器（encode_query），embedding 不可用时一并跳过，
    # 否则离线环境下自测会在这里抛 ConnectionRefused
        store2 = Store(idx_dir).load()
        ret = Retriever(store2, emb, cfg.load())
        hits = ret.search("6061 的时效温度和屈服强度", top_k=3)
        check("检索返回结果", len(hits) > 0, f"{len(hits)} 条")
        joined = " ".join(h["chunk"].text for h in hits)
        check("命中关键信息（175 C / 276 MPa）", ("175" in joined) or ("276" in joined),
              f"片段含 175={'175' in joined} 276={'276' in joined}")


# ---------------------------------------------------------------- 6. 抽取校验（模拟模型输出）

def t_extract_verify():
    print("\n[6] 抽取结果校验（用构造的模型输出来验证证据链，不调用 LLM）")
    text = ("Artificial ageing was performed at 175 C for 8 h. "
            "The yield strength reached 276 MPa.")
    chunk = Chunk.make("doc1", "test_paper.pdf", 1, "text", text)

    class FakeLLM:
        def chat_json(self, user, system="", retry_hint="", response_format=None):
            return {"records": [{
                "fields": {
                    "alloy": {"raw": "6061", "quote": "The yield strength reached 276 MPa."},
                    "process.aging_temp": {"raw": "175", "quote": "ageing was performed at 175 C"},
                    "process.aging_time": {"raw": "8", "quote": "for 8 h."},
                    "properties.YS": {"raw": "276", "quote": "yield strength reached 276 MPa"},
                    # 下面这条是编造的，必须被拦下
                    "properties.elongation": {"raw": "12", "quote": "elongation of 12 %"},
                },
                "note": "",
            }]}

    specs = build_schema_from_headers(
        ["合金牌号", "时效温度", "时效时间", "屈服强度", "延伸率"])["specs"]
    ext = Extractor(FakeLLM(), cfg.load())
    recs = ext.extract([chunk], specs)
    check("产出记录", len(recs) == 1)
    if recs:
        rec = recs[0]
        check("合金字段被识别", rec.alloy == "6061", f"实得 {rec.alloy}")
        ys = rec.properties.get("YS")
        check("屈服强度写入", ys is not None and ys.value == 276, f"实得 {ys.value if ys else None}")
        check("屈服强度证据已核验", bool(ys and ys.evidence and ys.evidence.verified))
        el = rec.properties.get("elongation")
        check("编造字段被标记为未核验",
              bool(el and el.evidence and not el.evidence.verified),
              "elongation 的引文不在原文中")
        check("记录带出警信息", len(rec.warnings) >= 1, f"{rec.warnings}")

    # 负值成分：实测踩过——模型把热力学表（ΔGseg，kJ/mol，含负值）的数值
    # 塞进成分字段，引文逐字命中、证据校验通过。这类值必须整字段跳过，
    # 且不能影响同一行里其他正常字段。
    tbl = Chunk.make("doc1", "test_paper.pdf", 8, "table",
                     "[表格] Table 2 The ΔGseg\n| | ΔGseg Gb of Mg | Si |\n|---|---|---|\n"
                     "| AC-AlMgSi | −9.127 | 1.2 |")

    class NegLLM:
        def chat_json(self, user, system="", retry_hint="", response_format=None):
            return {"records": [{
                "row_label": "AC-AlMgSi",
                "alloy": {"raw": "AC-AlMgSi", "quote": "AC-AlMgSi"},
                "fields": {
                    "composition.Mg": {"raw": "-9.127", "quote": "−9.127 | 1.2"},
                    "composition.Si": {"raw": "1.2", "quote": "−9.127 | 1.2"},
                },
                "note": ""}]}

    specs_neg = build_schema_from_headers(["合金牌号", "Mg", "Si"])["specs"]
    recs_neg = Extractor(NegLLM(), cfg.load()).extract([tbl], specs_neg)
    check("含不合理字段的记录仍被保留", len(recs_neg) == 1, f"实得 {len(recs_neg)} 条")
    if recs_neg:
        rec_neg = recs_neg[0]
        check("负值成分被跳过（不生成该字段）",
              "Mg" not in rec_neg.composition, f"实得字段 {list(rec_neg.composition)}")
        check("同一条记录里的正常字段不受影响",
              rec_neg.composition.get("Si") is not None,
              f"实得字段 {list(rec_neg.composition)}")
        # 负值成分现在会被"对齐校验"（表头是 ΔGseg Gb of Mg，不是 Mg 这一列）
        # 或"合理性校验"拦下，两者都留下了可追溯的说明文字。
        check("跳过/留空不合理值时留下说明",
              any("超出 0~100" in w or "未识别" in w for w in rec_neg.warnings),
              f"{rec_neg.warnings}")


def t_merge_records():
    print("\n[6b] 跨块记录合并（成分表 + 性能表应拼齐，不同条件不能误并）")
    from rag.extract.extractor import _dedupe
    from rag.models import Evidence, FieldValue, MaterialRecord

    def mk(alloy, comp=None, props=None, doc="d1"):
        r = MaterialRecord(record_id="", source_doc=doc, source_page=1)
        r.alloy = alloy
        for k, v in (comp or {}).items():
            r.composition[k] = FieldValue(
                value=v, unit="wt.%", evidence=Evidence(quote="c", verified=True))
        for k, v in (props or {}).items():
            r.properties[k] = FieldValue(
                value=v, unit="MPa", evidence=Evidence(quote="p", verified=True))
        return r

    merged = _dedupe([mk("A1", comp={"Si": 1.17}), mk("A1", props={"UTS": 408.7})])
    check("同牌号跨块合并成一条", len(merged) == 1, f"实得 {len(merged)} 条")
    if merged:
        r = merged[0]
        check("合并后成分与性能都在",
              r.composition.get("Si") is not None and r.properties.get("UTS") is not None,
              f"成分={list(r.composition)} 性能={list(r.properties)}")

    two = _dedupe([mk("A1", props={"UTS": 300}), mk("A1", props={"UTS": 400})])
    check("同牌号不同性能值不误并（不同工艺条件）", len(two) == 2, f"实得 {len(two)} 条")

    dup = _dedupe([mk("A1", comp={"Si": 1.17}), mk("A1", comp={"Si": 1.17})])
    check("完全重复去重成一条", len(dup) == 1, f"实得 {len(dup)} 条")

    noid = _dedupe([mk("", comp={"Si": 1.0}), mk("", props={"UTS": 200})])
    check("无牌号记录不参与合并（身份不明）", len(noid) == 2, f"实得 {len(noid)} 条")

    diff = _dedupe([mk("A1", comp={"Si": 1.0}, doc="d1"),
                    mk("A1", props={"UTS": 200}, doc="d2")])
    check("不同文献同牌号不合并", len(diff) == 2, f"实得 {len(diff)} 条")


def t_table_context():
    print("\n[6c] 表格参考上下文（同页正文只用于理解，不改块、标注为不可抽取）")
    table = Chunk.make("d1", "p.pdf", 3, "table",
                       "[表格] Table 2\n| No. | Si |\n|---|---|\n| A1 | 1.17 |")
    text_chunk = Chunk.make("d1", "p.pdf", 3, "text", "时效工艺为 175 °C/8 h。")
    seen: dict = {}

    class CapLLM:
        def chat_json(self, user, system="", retry_hint="", response_format=None):
            seen["user"] = user
            seen["rf"] = response_format
            return {"records": []}

    specs = build_schema_from_headers(["合金牌号", "Si"])["specs"]
    ext = Extractor(CapLLM(), cfg.load())

    ext.extract([table], specs,
                context_of=lambda c: "时效工艺为 175 °C/8 h，成分单位为 wt.%")
    check("表格块注入了参考上下文", "175 °C/8 h" in seen.get("user", ""))
    check("参考上下文标注「仅供理解」", "仅供理解" in seen.get("user", ""))
    check("抽取走结构化输出（json_object）", seen.get("rf") == {"type": "json_object"})

    seen.clear()
    ext.extract([text_chunk], specs, context_of=lambda c: "")
    check("正文块不注入参考上下文（只有目标块本身）",
          "参考上下文" not in seen.get("user", ""))


# ---------------------------------------------------------------- 7. Excel 写回与 diff

def t_excel(tmp: str):
    print("\n[7] Excel 写回：备份 + 追加 + 逐格 diff")
    xl = os.path.join(tmp, "target.xlsx")
    # 关键：把 output_dir 指向临时目录，否则自测会把备份写进项目的 data/output（污染生产目录）
    conf = copy.deepcopy(cfg.load())
    conf["project"]["output_dir"] = tmp
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "data"
    ws.append(["合金牌号", "时效温度", "时效时间", "屈服强度", "延伸率"])
    ws.append(["6063-T5", 175, 8, 214, 12])     # 已有一行，必须不被破坏
    wb.save(xl)

    target = ExcelTarget(xl, conf)
    heads = target.headers()
    check("读取表头", heads[0] == "合金牌号", f"{heads}")

    spec = build_schema_from_headers(heads)["specs"]
    ev = Evidence(quote="aged at 175 C for 8 h", chunk_id="c1", page=1,
                  doc_name="test_paper.pdf", verified=True)
    rec = MaterialRecord(alloy="6061-T6", record_id="r1",
                         source_doc="test_paper.pdf", source_page=1)
    rec.process["aging_temp"] = FieldValue(175, "°C", ev, {"kind": "single", "value": 175})
    rec.process["aging_time"] = FieldValue(8, "h", ev, {"kind": "single", "value": 8})
    rec.properties["YS"] = FieldValue(276, "MPa", ev, {"kind": "single", "value": 276})
    rec.properties["elongation"] = FieldValue(999, "%", Evidence(quote="不存在", verified=False),
                                              {"kind": "single", "value": 999})

    rows, evidence = ExcelTarget.to_rows([rec], spec, trusted_only=True)
    check("未核验字段被排除在数据行之外", "延伸率" not in (rows[0] if rows else {}),
          f"行内容 {rows[0] if rows else None}")
    check("证据明细覆盖全部字段（含未核验）", len(evidence) >= 4, f"{len(evidence)} 条")

    before = diffmod.snapshot(xl)
    w = target.write(rows, dry_run=False)
    check("写入成功且无非预期变更",
          w["unexpected_count"] == 0, f"unexpected={w['unexpected_count']} 全量={w['total_changed']}")
    check("备份文件已生成", os.path.exists(w["backup"]), w["backup"])
    check("原有行未被破坏",
          diffmod.snapshot(xl)["data"]["cells"].get((2, 1)) == "6063-T5")

    after = diffmod.snapshot(xl)
    d = diffmod.diff(before, after)
    check("diff 能报出新增单元格", d["added"] >= 4, f"added={d['added']}")
    print("\n  ---- diff 报告样例 ----")
    print("  " + diffmod.format_report(w).replace("\n", "\n  "))


# ---------------------------------------------------------------- 8. 终端渲染

def t_format_text():
    print("\n[8] 默认行为：不指定 Excel 时，结果直接渲染在终端（不碰任何文件）")
    from rag.cli import format_records_text
    from rag.extract.schema import FIELD_DICT

    ev = Evidence(quote="aged at 175 C for 8 h", chunk_id="c1", page=1,
                  doc_name="t.pdf", verified=True)
    rec = MaterialRecord(alloy="6061-T6", record_id="r1",
                         source_doc="t.pdf", source_page=1)
    rec.properties["YS"] = FieldValue(276, "MPa", ev, {"kind": "single", "value": 276})
    rec.process["aging_temp"] = FieldValue(175, "°C", ev, {"kind": "single", "value": 175})
    rec.properties["elongation"] = FieldValue(
        999, "%", Evidence(quote="编造的引文", verified=False), {"kind": "single", "value": 999})

    specs = {f.key: f for f in FIELD_DICT}
    text = format_records_text([rec], specs)
    check("渲染出合金牌号", "6061-T6" in text)
    check("渲染出字段中文名（未指定表头时用词典名）", "屈服强度" in text)
    check("未通过核验的值不显示", "999" not in text)
    check("统计了有效字段数", "个有效字段" in text)
    print("\n  ---- 终端渲染样例 ----")
    print("  " + text.replace("\n", "\n  "))


# ---------------------------------------------------------------- main

def t_rules():
    """第 9 组：行标签过滤 / 预测值字段 / 单位契约。"""
    print("\n[9] 行标签过滤、预测值字段、单位契约")
    from rag.extract.extractor import _is_bad_label
    for lab, want in [("Step sizes 5", True), ("Limit (wt. %)", True),
                      ("Ranges 500–570", True), ("Mean", True),
                      ("A1", False), ("6013", False), ("AMSY-Exp", False),
                      ("Mg", False)]:
        check(f"行标签 {lab!r} -> {'过滤' if want else '保留'}",
              _is_bad_label(lab) == want)
    from rag.extract.schema import BY_KEY
    check("存在 UTS 预测值字段", "properties.UTS_pred" in BY_KEY)
    check("UTS 预测值单位 MPa", BY_KEY["properties.UTS_pred"].unit == "MPa")
    check("延伸率为实测字段且单位为 %",
          BY_KEY["properties.elongation"].unit == "%")
    check("延伸率预测值单位为 %",
          BY_KEY["properties.elongation_pred"].unit == "%")
    check("裸 % 不再被归一成 wt.%", canonical_unit("%") == "%")
    from rag.extract.extractor import _pred_only_caption
    check("表题『Predicted performance...』判为纯预测表",
          _pred_only_caption("[表格] Table 2 Predicted performance and the corresponding composition"))
    check("表题『Predicted and experimental...』不算纯预测表",
          not _pred_only_caption("[表格] Table 5 Predicted and experimental composition comparison"))
    check("表题『Theoretical vs actual...』不算纯预测表",
          not _pred_only_caption("[表格] Table 5 Theoretical component performance vs. actual component performance"))


def t_schemas(tmp: str):
    """第 10 组：动态 Schema（重构回归基线 + 校验 + 优先级）。"""
    print("\n[10] 动态 Schema：内置回归、模板校验、Schema 优先级")
    from rag.extract.schema import FIELD_DICT, schema_brief
    from rag.schemas import builtin_schema, from_text

    # ---- 重构回归基线：内置模板必须与旧词典逐字符一致 ----
    old = schema_brief({f.key: f for f in FIELD_DICT})
    new = schema_brief(builtin_schema().to_specs())
    check("内置模板字段数 = 94", len(builtin_schema().to_specs()) == 94,
          f"实得 {len(builtin_schema().to_specs())}")
    check("内置模板渲染与旧词典逐字符一致", old == new,
          f"旧 {len(old)} 字符 / 新 {len(new)} 字符")

    # ---- 合法模板 ----
    good = """
name: 测试模板
key_field: alloy
groups:
  - key: material
    label: 材料
    bucket: meta
    fields:
      - {key: alloy, label: 合金牌号, type: text, required: true}
  - key: chemical
    label: 化学成分
    fields:
      - {key: Si, label: Si, type: number, unit: wt.%}
      - {key: Fe, label: Fe, type: number, unit: wt.%}
  - key: props
    label: 力学性能
    fields:
      - {key: UTS, label: 抗拉强度, type: number, unit: MPa}
"""
    s, errs, warns = from_text(good)
    check("合法模板解析通过", s is not None and not errs, f"{errs}")

    # ---- 显示名（YAML 里的 name）与文件名不一致时也能解析 ----
    # 场景：界面模板列表显示「铝合金精简模板」，但文件名是 example_al_alloy.yaml。
    # 界面点选时传的是显示名，resolve 必须能兜底按 name 字段扫描。
    # 把 schemas_dir 临时指向 tmp，避免污染 data/schemas。
    import rag.schemas.loader as _L
    _orig_dir = _L.schemas_dir
    _L.schemas_dir = lambda: tmp
    try:
        _p = os.path.join(tmp, "alias_test.yaml")
        with open(_p, "w", encoding="utf-8") as f:
            f.write(good.replace("name: 测试模板", "name: 我的中文模板"))
        s_by_disp, e_disp, _w = _L.resolve("我的中文模板")
        check("按显示名解析（显示名≠文件名）",
              s_by_disp is not None and s_by_disp.name == "我的中文模板",
              f"{e_disp[0]['msg'] if e_disp else ''}")
        s_by_file, _e2, _w2 = _L.resolve("alias_test")
        check("按文件名解析仍然有效", s_by_file is not None and s_by_file.name == "我的中文模板")
        s_miss, e_miss, _w3 = _L.resolve("根本不存在")
        check("不存在的模板正确报错",
              s_miss is None and "找不到" in (e_miss[0]["msg"] if e_miss else ""))
    finally:
        _L.schemas_dir = _orig_dir

    # ---- 字段拖拽排序的后端：move 操作 ----
    from rag.schemas.loader import apply_field_edit as _afe
    _mv = """
name: mv
key_field: alloy
groups:
  - key: m
    label: 材料
    bucket: meta
    fields:
      - {key: alloy, label: 合金牌号, type: text}
  - key: chem
    label: 成分
    bucket: composition
    fields:
      - {key: Si, label: Si, type: number, unit: wt.%}
      - {key: Fe, label: Fe, type: number, unit: wt.%}
      - {key: Mg, label: Mg, type: number, unit: wt.%}
  - key: prop
    label: 性能
    bucket: properties
    fields:
      - {key: UTS, label: 抗拉强度, type: number, unit: MPa}
"""
    _seq = lambda s, k: [f.key for g in s.groups if g.key == k for f in g.fields]

    t1, _e, _i = _afe(_mv, op="move", group="chem",
                      keys=["composition.Fe"], before="composition.Si")
    s1 = from_text(t1)[0]
    check("move 同组排序", _seq(s1, "chem") == ["Fe", "Si", "Mg"], str(_seq(s1, "chem")))

    t2, _e, _i = _afe(_mv, op="move", group="chem",
                      keys=["properties.UTS"], before="composition.Mg")
    s2 = from_text(t2)[0]
    check("move 跨组移动", _seq(s2, "chem") == ["Si", "Fe", "UTS", "Mg"])
    check("move 后空组被移除", all(g.key != "prop" for g in s2.groups))

    t3, _e, _i = _afe(_mv, op="move", group="chem", keys=["composition.Si"], before="")
    s3 = from_text(t3)[0]
    check("move 无 before 时追加到组末尾", _seq(s3, "chem") == ["Fe", "Mg", "Si"])

    t4, e4, _i = _afe(_mv, op="move", group="chem", keys=["composition.Xx"], before="")
    check("move 不存在的字段报错", t4 is None and "找不到" in e4[0]["msg"])

    t5, e5, _i = _afe(_mv, op="move", group="nope", keys=["composition.Si"], before="")
    check("move 不存在的组报错", t5 is None and "不存在" in e5[0]["msg"])

    t6, _e, _i = _afe(_mv, op="add", group="chem", fields=["Ni"])
    t7, _e, _i = _afe(t6, op="remove", keys=["composition.Ni"])
    check("add/remove 在加入 move 后回归正常", t7 is not None and "Ni" not in t7)

    if s:
        specs = s.to_specs()
        check("内部字段名用「桶.字段」", "composition.Si" in specs and "properties.UTS" in specs,
              f"{list(specs)}")
        check("meta 组字段用裸名", "alloy" in specs, f"{list(specs)}")
        check("分组 bucket 自动推断（chemical -> composition）",
              s.groups[1].resolved_bucket() == "composition",
              s.groups[1].resolved_bucket())
        check("schema_hash 稳定且非空",
              s.schema_hash and s.schema_hash == from_text(good)[0].schema_hash, s.schema_hash)

    # ---- 非法模板：必须逐条报错 ----
    bad_type = good.replace("type: number, unit: wt.%", "type: 乱写", 1)
    s2, e2, _ = from_text(bad_type)
    check("非法 type 被拒绝", s2 is None and any("type" in x["msg"] for x in e2), f"{e2}")

    bad_enum = good + "      - {key: temper, label: 状态, type: enum}\n"
    s3, e3, _ = from_text(bad_enum)
    check("enum 未给取值被拒绝", s3 is None and any("enum" in x["msg"] for x in e3), f"{e3}")

    dup = good.replace("key: Fe", "key: Si")
    s4, e4, _ = from_text(dup)
    check("字段内部名重复被拒绝", s4 is None and any("冲突" in x["msg"] for x in e4), f"{e4}")

    s5, e5, _ = from_text("name: x\ngroups: []")
    check("空 groups 被拒绝", s5 is None and bool(e5), f"{e5}")

    # ---- Schema 优先级：Excel 表头覆盖模板 ----
    import openpyxl
    from rag import config as cfg
    from rag.pipeline import RagPipeline
    xlsx = os.path.join(tmp, "target.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "data"
    ws.append(["合金牌号", "Si", "屈服强度"])
    wb.save(xlsx)

    p = RagPipeline(lazy=False)
    res = p.extract(excel_path=xlsx, schema=good, mode="table", doc_ids=["不存在的文献"])
    check("指定 Excel 时以表头为准（覆盖 Schema）",
          (res.get("schema") or {}).get("mode") == "excel_headers",
          str(res.get("schema", {}).get("mode")))
    check("表头覆盖后 specs 来自 Excel 列",
          any(s.column == "屈服强度" for s in res["specs"].values()),
          str([s.column for s in res["specs"].values()]))

    res2 = p.extract(schema=good, mode="table", doc_ids=["不存在的文献"])
    check("未指定 Excel 时用传入的 Schema",
          (res2.get("schema") or {}).get("mode") == "schema")
    res3 = p.extract(mode="table", doc_ids=["不存在的文献"])
    check("都没指定时回落到内置模板",
          (res3.get("schema") or {}).get("mode") == "builtin")

    # ---- 字段手输：解析与增删 ----
    from rag.schemas.loader import apply_field_edit, parse_field_spec
    sp = parse_field_spec("Si")
    check("手输元素符号自动带出类型/单位",
          sp.get("key") == "Si" and sp.get("type") == "number" and sp.get("unit") == "wt.%",
          str(sp))
    sp = parse_field_spec("抗拉强度")
    check("手输中文名映射到内置字段",
          sp.get("key") == "UTS" and sp.get("unit") == "MPa", str(sp))
    sp = parse_field_spec("自定义指标:number:%")
    check("手输「名称:类型:单位」生效",
          sp.get("key") == "自定义指标" and sp.get("type") == "number" and sp.get("unit") == "%",
          str(sp))
    check("手输非法类型被拒", bool(parse_field_spec("X:乱写").get("error")),
          str(parse_field_spec("X:乱写")))

    base = ("name: t\nkey_field: alloy\ngroups:\n"
            "  - {key: material, label: 材料, bucket: meta, fields: [{key: alloy, type: text}]}\n"
            "  - {key: props, label: 性能, bucket: properties, fields: [{key: UTS, type: number}]}\n")
    txt, errs, info = apply_field_edit(base, "add", "props", ["Ni, 抗拉强度", "自定义指标:number:%"])
    check("加字段：Ni 与自定义指标成功，抗拉强度因已存在被跳过",
          not errs and info["added"] == ["Ni", "自定义指标"] and info["skipped"] == ["UTS"],
          f"{info}")
    check("加字段后 YAML 含新键", txt and "key: Ni" in txt and "自定义指标" in txt)
    txt2, errs2, info2 = apply_field_edit(txt, "remove", keys=["properties.Ni"])
    check("删字段生效", not errs2 and "properties.Ni" in info2["removed"]
          and "key: Ni" not in (txt2 or ""), str(info2))
    _t, errs3, _i = apply_field_edit(base, "add", "不存在的组", ["X"])
    check("目标分组不存在时报错", bool(errs3), str(errs3))
    _t2, errs4, _i2 = apply_field_edit(
        "name: t\ngroups:\n  - {key: g, label: G, bucket: properties, fields: [{key: A}]}\n",
        "remove", keys=["properties.A"])
    check("删空最后一个字段被拒绝", bool(errs4), str(errs4))


def t_clean():
    """第 11 组：导出前的基础清洗（只改字符层）。"""
    print("\n[11] 导出清洗：去空格 / 修数字内空格 / 全半角 / 去重复单位")
    from rag.writeback.clean import clean_value, clean_rows, dedupe_rows
    R = ["unicode", "trim", "space_in_number", "strip_unit"]

    got, used = clean_value("  0.  62 ", R, "wt.%")
    check("去空格 + 修数字内空格", got == "0.62", f"{got!r} {used}")

    got, _ = clean_value("８０．５", R, "MPa")
    check("全角数字转半角", got == "80.5", repr(got))

    got, used = clean_value("Fig. 5", R, "")
    check("非数值串不粘空格（Fig. 5 保持不变）", got == "Fig. 5" and not used, f"{got!r} {used}")

    got, _ = clean_value("0.62 wt.%", R, "wt.%")
    check("值尾重复单位可去掉", got == "0.62", repr(got))

    got, _ = clean_value("0–2.40", R, "wt.%")
    check("区间写法不被破坏", got == "0–2.40", repr(got))

    got, _ = clean_value("380 ± 2", R, "MPa")
    check("mean±sd 不被拆开", got == "380 ± 2", repr(got))

    rows = [{"a": "0.62", "b": "x"}, {"a": "0.62", "b": "x"}, {"a": "0.63", "b": "x"}]
    deduped, removed = dedupe_rows(rows)
    check("完全重复行被去掉", len(deduped) == 2 and removed == 1, f"{len(deduped)} / {removed}")

    out, rep = clean_rows([{"Si": " 0.62 ", "Fe": "0.19"}], ["trim"], {"Si": "wt.%"})
    check("清洗报告统计正确", rep["changed_cells"] == 1 and rep["by_rule"].get("trim") == 1,
          str(rep))

    # ---- 第 13 组：物理合理性校验（抓"引文存在但语义错误"的幻觉填充）----
    from rag.extract.normalize import sanity_issue, composition_sum

    check("Al=20 wt.% 被判为异常（铝合金基体）",
          "偏低" in sanity_issue("composition.Al", 20, "wt.%"))
    check("Al=40.5 wt.% 被判为异常",
          "偏低" in sanity_issue("composition.Al", 40.5, "wt.%"))
    check("Al=92.3 wt.% 正常不报", sanity_issue("composition.Al", 92.3, "wt.%") == "")
    check("Al='Bal.'（余量）不报", sanity_issue("composition.Al", "Bal.", "wt.%") == "")
    check("Si=20 wt.% 不报（过共晶 Al-Si 合理）",
          sanity_issue("composition.Si", 20, "wt.%") == "")
    check("成分 >100 被判异常", "超出" in sanity_issue("composition.Al", 150, "wt.%"))
    check("非 wt.% 单位不做成分校验",
          sanity_issue("composition.Al", 20, "at.%") == "")
    check("非成分字段不做该校验",
          sanity_issue("properties.UTS", 20, "MPa") == "")

    class _FV:
        def __init__(self, v, u): self.value, self.unit = v, u
    s_ok = composition_sum({"Si": _FV(1.17, "wt.%"), "Mg": _FV(1.01, "wt.%")})
    check("成分合计计算正确", abs((s_ok or 0) - 2.18) < 1e-6, f"{s_ok}")
    check("空成分返回 None", composition_sum({}) is None)
    check("非 wt.% 不纳入合计", composition_sum({"UTS": _FV(400, "MPa")}) is None)


def t_chunk_plan(tmp: str):
    """抽取规模控制：模式差异、auto 的「数值+单位」判定、块数上限、排序。"""
    print("\n[12] 抽取规模：块筛选 / 模式差异 / 块数上限")
    from rag.pipeline import RagPipeline

    idx = os.path.join(tmp, "idx_plan")
    store = Store(idx)
    cs = [
        Chunk.make("d1", "a.pdf", 1, "table", "Table 1 Si 0.62 wt.% Mg 0.90 wt.%"),
        Chunk.make("d1", "a.pdf", 2, "table", "Table 2 UTS 320 MPa YS 276 MPa"),
        # 只有关键词、没有「数值+单位」——不应被 auto 选中
        Chunk.make("d1", "a.pdf", 3, "text",
                   "The aged alloy shows excellent strength and good elongation."),
        Chunk.make("d1", "a.pdf", 4, "text", "Yield strength reached 276 MPa after 8 h."),
        Chunk.make("d2", "b.pdf", 1, "table", "Table 1 Cu 0.50 wt.%"),
    ]
    emb = Embedder(cfg.load())
    emb.load()
    # 这里只是要建一个可用的 store 来测块筛选，向量内容无关紧要；
    # embedding 服务不可用时用占位向量，保证自测能离线跑。
    try:
        _vecs = emb.encode_documents([c.text for c in cs])
    except Exception:
        import numpy as np
        _vecs = np.zeros((len(cs), 8), dtype=np.float32)
    store.add(cs, _vecs)
    store.save()

    p = RagPipeline()
    p.store = Store(idx).load()

    t = p.pick_chunks(mode="table")
    check("table 模式只要表格块", len(t) == 3 and all(c.kind == "table" for c in t), f"{len(t)} 块")

    a = p.pick_chunks(mode="auto")
    kinds = [c.kind for c in a]
    check("auto = 表格 + 含数值单位的段落", len(a) == 4, f"{len(a)} 块 {kinds}")
    check("auto 排除了只有关键词的段落",
          not any(c.page == 3 and c.kind == "text" for c in a), "p3 不应入选")
    check("auto 收下了 276 MPa 那段", any(c.page == 4 for c in a))

    check("all 模式不筛选", len(p.pick_chunks(mode="all")) == 5)

    lim = p.pick_chunks(mode="table", max_chunks=2)
    check("max_chunks 截断生效", len(lim) == 2, f"{len(lim)} 块")
    check("截断前表格优先、按页码排",
          [c.page for c in lim] == [1, 2], str([c.page for c in lim]))

    # 继续抽取：skip 跳过排序后的前 N 块，从上次完成位置接着抽
    full_ids = [c.chunk_id for c in p.pick_chunks(mode="table")]
    skipped = p.pick_chunks(mode="table", skip=1)
    check("skip 跳过前 N 块", len(skipped) == 2 and skipped[0].chunk_id == full_ids[1],
          f"{len(skipped)} 块")
    combo = p.pick_chunks(mode="table", skip=1, max_chunks=1)
    check("skip 与 max_chunks 组合", len(combo) == 1 and combo[0].chunk_id == full_ids[1])

    one = p.pick_chunks(mode="table", doc_ids=["d2"])
    check("doc_ids 只取指定文献", len(one) == 1 and one[0].doc_id == "d2")

    counts = {m: len(p.pick_chunks(mode=m)) for m in ("table", "auto", "all")}
    check("三模式计数与 plan 接口一致",
          counts == {"table": 3, "auto": 4, "all": 5}, str(counts))

    # 按块 id 续抽：切换抽取范围后，块列表会变，靠"跳过前 N 块"会指错位置。
    # 改成排除已抽 id 后，换模式续抽也不该抽到重复块。
    _b1 = p.pick_chunks(mode="table", max_chunks=3)
    _ids1 = [c.chunk_id for c in _b1]
    _b2 = p.pick_chunks(mode="all", max_chunks=3, exclude_ids=_ids1)
    _ids2 = [c.chunk_id for c in _b2]
    check("切换范围后按 id 续抽不重复",
          not (set(_ids1) & set(_ids2)), f"重叠 {len(set(_ids1) & set(_ids2))} 个")
    check("排除已抽块后取到的是剩余块",
          len(_ids2) == counts["all"] - len(_ids1),
          f"实得 {len(_ids2)}，预期 {counts['all'] - len(_ids1)}")

    # 服务生命周期配置必须由接口下发（界面文案不许写死秒数——
    # 曾把 "8 秒" 硬编码在模板里，配置改成 15 秒后界面仍在骗人）
    _st = p.status().get("web") or {}
    _grace = _st.get("shutdown_grace_s")
    _cfg_grace = float(cfg.get("web.shutdown_grace_s", 15))
    check("status 下发 web.shutdown_grace_s 且与配置一致",
          isinstance(_grace, (int, float)) and abs(float(_grace) - _cfg_grace) < 1e-6,
          f"接口={_grace} 配置={_cfg_grace}")
    check("status 下发 shutdown_on_page_close 布尔值",
          isinstance(_st.get("shutdown_on_page_close"), bool),
          f"实得 {_st.get('shutdown_on_page_close')!r}")

    # 中途停止：界面点「停止」后不该把剩下的块继续跑完
    calls = {"n": 0}

    class CountLLM:
        def chat_json(self, user, system="", retry_hint="", response_format=None):
            calls["n"] += 1
            return {"records": []}

    specs = build_schema_from_headers(["合金牌号"])["specs"]
    cs5 = [Chunk.make("d1", "a.pdf", i, "text", "nothing here") for i in range(1, 6)]
    _ext_stop = Extractor(CountLLM(), cfg.load())
    _ext_stop.extract(cs5, specs, workers=1, should_stop=lambda: calls["n"] >= 2)
    check("中途停止：只跑完已开始的批次", calls["n"] == 2, f"调了 {calls['n']} 次 / 共 5 块")
    check("停止后 stopped_early=True（供上层如实报告）", _ext_stop.stopped_early is True)

    calls["n"] = 0
    _ext_ok = Extractor(CountLLM(), cfg.load())
    _ext_ok.extract(cs5, specs, workers=1)
    check("不停止时全部跑完", calls["n"] == 5, f"调了 {calls['n']} 次")
    check("未停止时 stopped_early=False", _ext_ok.stopped_early is False)

    # ---- 索引一致性：块与向量必须同步 ----
    import numpy as _np
    from rag.index.store import Store as _Store
    _sd = os.path.join(tmp, "store_sync")
    _st = _Store(_sd)
    _c = Chunk.make("dx", "x.pdf", 1, "text", "hello")
    _st.add([_c], _np.zeros((1, 4), dtype=_np.float32))
    _raised = False
    try:
        _st.add([_c], None)          # 有向量时再塞无向量的块 -> 下标会错位
    except ValueError:
        _raised = True
    check("块与向量不同步时拒绝写入（防检索返回错误块）", _raised)

    # ---- 模型不可达时，健康检查必须快速返回 ----
    # 曾踩过：本地端口未监听时 HTTP 连接要等约 2 秒，界面状态接口被拖住，
    # 启动脚本轮询（超时 1 秒）每次都失败，表现成"双击 start_web.vbs 没反应"。
    import time as _time
    from rag.llm.client import LLMClient as _LLM
    _probe_client = _LLM()
    _probe_client.base_url = "http://127.0.0.1:19999/v1"     # 故意不可达
    _t0 = _time.time()
    _h = _probe_client.health()
    _dt = _time.time() - _t0
    check("模型不可达时 health 快速返回（<1 秒）",
          (not _h.get("ok")) and _dt < 1.0, f"{_dt:.2f}s, ok={_h.get('ok')}")

    # ---- 刷新后接管任务：_running_job 的状态判定 ----
    # _RESTORED_DONE=True 屏蔽"从磁盘恢复上一批"这一步：本用例只测内存里的接管逻辑，
    # 不能让本机 data/output/last_extract.json 里真实的上批结果掺进来。
    import rag.web.app as _WA
    from rag.web.app import _running_job as _rj, _JOBS as _J
    _saved = dict(_J)
    _saved_restored, _WA._RESTORED_DONE = _WA._RESTORED_DONE, True
    try:
        _J.clear()
        check("无任务时不接管", _rj() == {})
        _J["run1"] = {"q": None, "done": False, "result": None, "error": None,
                      "started_at": 100.0, "last": {"done": 3, "total": 10},
                      "req": {"skip": 5, "max_chunks": 10}}
        r = _rj()
        check("接管运行中的任务",
              r.get("job_id") == "run1" and r.get("done") is False
              and r.get("progress", {}).get("done") == 3
              and r.get("req", {}).get("skip") == 5)
        _J["run1"]["done"] = True
        _J["run1"]["result"] = {"records_count": 7}
        _J["run1"]["finished_at"] = _time.time()   # 结果只在时限内交还（web.adopt_ttl_s）
        r2 = _rj()
        check("接管已跑完未收取的结果",
              r2.get("done") is True and r2.get("result", {}).get("records_count") == 7)
        _J.clear()
        _J["err1"] = {"q": None, "done": True, "result": None, "error": "boom",
                      "started_at": 1.0, "last": {}, "req": {}}
        check("出错的任务不被接管", _rj() == {})
        _J.clear()
        _J["old"] = {"q": None, "done": True, "result": {"records_count": 1},
                     "error": None, "started_at": 1.0, "last": {}, "req": {}}
        _J["new"] = {"q": None, "done": False, "result": None, "error": None,
                     "started_at": 2.0, "last": {}, "req": {}}
        check("优先接管正在跑的任务", _rj().get("job_id") == "new")
    finally:
        _J.clear()
        _J.update(_saved)
        _WA._RESTORED_DONE = _saved_restored


# ---------------------------------------------------------------- 15. 启动脚本

def t_launcher() -> None:
    """start_web.vbs 必须保持纯 ASCII。

    Windows Script Host 按系统 ANSI 代码页（cp936/GBK）解析 .vbs，不认 UTF-8。
    UTF-8 中文被当成 GBK 双字节，奇数个汉字会留下一个悬空前导字节，把字符串的
    收尾引号吞掉 → 800A0409「未结束的字符串常量」，脚本一行都执行不到。
    """
    print("\n[15] 启动脚本")
    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    p = os.path.join(root, "start_web.vbs")
    check("start_web.vbs 存在", os.path.isfile(p))
    raw = open(p, "rb").read()
    bad = [i for i, b in enumerate(raw) if b > 127]
    check("start_web.vbs 为纯 ASCII（写中文会让 WSH 解析失败）", not bad,
          f"非 ASCII 字节 {len(bad)} 个" if bad else "")
    lines = raw.decode("ascii", "replace").splitlines()
    odd = [i for i, ln in enumerate(lines, 1) if ln.count('"') % 2]
    check("start_web.vbs 每行引号成对", not odd, f"行号 {odd}" if odd else "")
    check("start_web.vbs 用文件标记判定就绪", "web.ready" in raw.decode("ascii", "replace"))


# ---------------------------------------------------------------- 16. 抽取事件订阅

def t_job_stream() -> None:
    """抽取进度事件必须广播给所有订阅者，且留下日志供后连者补发。

    早先每个任务只有一条 Queue：刷新页面后新订阅会和旧连接抢同一条队列，
    进度事件被分掉，chunk_ids 到不了新页面 → 「继续抽取」的按钮消失。
    """
    print("\n[16] 抽取事件的订阅与补发")
    import threading

    from rag.web import app as webapp

    job = {"events": [], "cv": threading.Condition(), "done": False,
           "result": None, "error": None, "started_at": 0.0,
           "last": {}, "done_ids": [], "req": {}}

    def consume(sink: list) -> None:
        for ev in webapp._iter_events(job):
            if ev is None:                  # 心跳
                continue
            sink.append(ev)
            if ev.get("type") in ("done", "error"):
                break

    early: list = []
    t1 = threading.Thread(target=consume, args=(early,), daemon=True)
    t1.start()
    time.sleep(0.05)
    webapp._publish_event(job, {"type": "progress", "done": 1, "chunk_ids": ["c1"]})

    late: list = []                         # 半路才连上的订阅者
    t2 = threading.Thread(target=consume, args=(late,), daemon=True)
    t2.start()
    time.sleep(0.05)
    webapp._publish_event(job, {"type": "progress", "done": 2, "chunk_ids": ["c2"]})
    webapp._publish_event(job, {"type": "done", "chunk_ids": ["c1", "c2"]})
    t1.join(5)
    t2.join(5)

    check("先连的订阅者收到全部 3 条事件", len(early) == 3, f"实得 {len(early)}")
    check("后连的订阅者补齐历史事件（不再被抢走）",
          [e.get("chunk_ids") for e in late] == [["c1"], ["c2"], ["c1", "c2"]],
          str(late))
    check("done 事件带本批完整 chunk_id（漏了进度也能补）",
          early[-1].get("chunk_ids") == ["c1", "c2"])
    check("事件不重复投送（日志+游标，一次一条）",
          [e.get("done") for e in early] == [1, 2, None])


# ---------------------------------------------------------------- 17. 任务接管

def t_adopt_ttl() -> None:
    """已完成任务的结果只在时限内交还，否则刷新每次都要搬运+重绘上一批结果。"""
    print("\n[17] 刷新接管的时效")
    import threading

    from rag.web import app as webapp

    saved = dict(webapp._JOBS)
    saved_restored, webapp._RESTORED_DONE = webapp._RESTORED_DONE, True
    now = 1_000_000.0
    try:
        webapp._JOBS.clear()
        webapp._JOBS["fresh"] = {"events": [], "cv": threading.Condition(),
                                 "done": True, "result": {"records_count": 1},
                                 "finished_at": now - 60, "started_at": now - 300,
                                 "last": {}, "done_ids": ["c1"], "req": {}}
        webapp._JOBS["old"] = {"events": [], "cv": threading.Condition(),
                               "done": True, "result": {"records_count": 1},
                               "finished_at": now - 7200, "started_at": now - 7300,
                               "last": {}, "done_ids": ["c2"], "req": {}}
        fresh = webapp._running_job(now)
        check("刚结束的任务：结果照常交还",
              fresh.get("done") and fresh.get("result") is not None)
        check("超时任务：结果不再交还（刷新不再搬运旧结果）",
              webapp._running_job(now + 700) is not None
              and webapp._running_job(now + 700).get("result") is None)
        check("超时任务仍还原「已抽到哪」",
              webapp._running_job(now + 700).get("done_ids") in (["c1"], ["c2"]))
        proc = webapp._JOBS  # 正在跑的任务不受时限影响
        proc["live"] = {"events": [], "cv": threading.Condition(), "done": False,
                        "result": None, "started_at": now, "last": {},
                        "done_ids": [], "req": {}}
        check("正在跑的任务永远可接管",
              webapp._running_job(now + 99999).get("job_id") == "live")
    finally:
        webapp._JOBS.clear()
        webapp._JOBS.update(saved)
        webapp._RESTORED_DONE = saved_restored


# ---------------------------------------------------------------- 18. 对齐 + 合理性

_TAB = (
    "[表格] Table 2 Chemical compositions of the alloys tested in this work, wt.%.\n"
    "| Alloy | Si | Mg | Cu | Fe | Al |\n"
    "|---|---|---|---|---|---|\n"
    "| A1 | 0.62 | 0.85 | 0.11 | 0.21 | Bal. |\n"
)
_TAB_NO_SI = (
    "[表格] Table 2 The ΔGseg of solute atoms at the grain boundary (kJ/mol).\n"
    "| Element | of Mg | of Si | of Cu |\n"
    "|---|---|---|---|\n"
    "| value | -9.127 | -13.628 | -6.234 |\n"
)


def t_alignment() -> None:
    """字段名必须能在表头/行首找到对应，否则判「未识别」——不许猜。"""
    print("\n[18] 字段-数值对齐")
    from rag.extract.align import alignment_issue
    from rag.extract.schema import FieldSpec

    si = FieldSpec(key="composition.Si", label="Si", group="composition",
                   aliases=["硅", "Silicon"], kind="number", unit="wt.%")
    mg = FieldSpec(key="composition.Mg", label="Mg", group="composition",
                   aliases=["镁", "Magnesium"], kind="number", unit="wt.%")

    check("表头有 Si 列 -> 可对齐", alignment_issue(si, "0.62", _TAB, "table") == "")
    check("表头有 Mg 列 -> 可对齐", alignment_issue(mg, "0.85", _TAB, "table") == "")
    bad = alignment_issue(si, "-13.628", _TAB_NO_SI, "table")
    check("表头只有 of Si（不是 Si 列）-> 判未识别", bool(bad), bad[:40])
    # 描述性长表头仍要认；能量/热力学列（ΔGseg）一律不认
    long_tbl = ("[表格] Chemical composition\n"
                "| Item | Mass fraction of Si | Mass fraction of Mg |\n|---|---|---|\n"
                "| A1 | 0.62 | 0.85 |\n")
    check("描述性表头 Mass fraction of Si -> 可对齐",
          alignment_issue(si, "0.62", long_tbl, "table") == "")
    # 元素写在行首（元素按行排的表）也要认
    row_tbl = ("[表格] Chemical composition (wt.%)\n| Item | A1 | A2 |\n|---|---|---|\n"
               "| Si | 0.62 | 0.71 |\n| Mg | 0.85 | 0.90 |\n")
    check("元素写在行首 -> 可对齐", alignment_issue(si, "0.62", row_tbl, "table") == "")
    # 正文块：引文上下文里要能找到字段写法
    prose = ("The alloy contains 0.62 wt.% Si and 0.85 wt.% Mg in this study. "
             + "The samples were homogenised at 550 C for 6 h and then water quenched. "
             + "Tensile tests were carried out at room temperature with a strain rate "
             + "of 1e-3 per second on an Instron machine, and the reported values are "
             + "the average of three parallel specimens measured in this study.")
    check("正文引文含字段名 -> 可对齐",
          alignment_issue(si, "0.62 wt.% Si", prose, "text") == "")
    check("正文引文上下文里没有字段名 -> 判未识别",
          bool(alignment_issue(si, "on an Instron machine", prose, "text")))


def t_sanity_range() -> None:
    """成分越界（>100 / 超出元素常见上限）必须能被抓到并可整条丢弃。"""
    print("\n[19] 合理性与一致性校验")
    from rag.extract.normalize import ELEMENT_MAX_WT, sanity_issue

    check("Si=0.62 wt.% 正常", sanity_issue("composition.Si", 0.62, "wt.%") == "")
    check("Si=93.5 wt.% 超上限（用户实测的错值）",
          "超出" in sanity_issue("composition.Si", 93.5, "wt.%", ELEMENT_MAX_WT))
    check("Mg=95 wt.% 超上限",
          "超出" in sanity_issue("composition.Mg", 95.0, "wt.%", ELEMENT_MAX_WT))
    check("Si=-9.127 wt.% 负值被抓",
          "超出 0~100" in sanity_issue("composition.Si", -9.127, "wt.%", ELEMENT_MAX_WT))
    check("非 wt.% 字段不受成分上限影响",
          sanity_issue("properties.UTS", 350.0, "MPa", ELEMENT_MAX_WT) == "")
    check("上限可覆盖（Si 放到 30）",
          sanity_issue("composition.Si", 28.0, "wt.%", {"Si": 30.0}) == "")

    # 记录级：成分合计超过上限 -> 整条丢弃（写进 dropped）
    from rag.extract.extractor import Extractor

    conf = {"extract": {"verify_quotes": True, "max_evidence_chars": 400,
                        "sanity": {"enabled": True, "on_violation": "drop",
                                   "sum_max_wt": 105, "on_sum_violation": "drop_row",
                                   "drop_row_if_violations": 3},
                        "alignment": {"enabled": False}}}
    ex = Extractor.__new__(Extractor)
    ex.conf = conf
    ex.dropped = []
    ex.stopped_early = False
    obj = {"records": [{"row_label": "A1", "fields": {
        "composition.Si": {"raw": "93.5", "quote": "93.5"},
        "composition.Mg": {"raw": "90.2", "quote": "90.2"},
    }}]}
    from rag.extract.schema import FieldSpec
    from rag.models import Chunk

    specs = {"composition.Si": FieldSpec(key="composition.Si", label="Si",
                                         group="composition", kind="number", unit="wt.%"),
             "composition.Mg": FieldSpec(key="composition.Mg", label="Mg",
                                         group="composition", kind="number", unit="wt.%")}
    ch = Chunk.make("d1", "x.pdf", 1, "text", "93.5 90.2")
    recs = ex._to_records(obj, [ch], specs)
    check("两个成分都越界 -> 整条丢弃", recs == [] and len(ex.dropped) == 1,
          str(ex.dropped[:1]))
    check("丢弃原因写清了合计/越界",
          bool(ex.dropped) and ("合计" in ex.dropped[0]["reason"]
                                or "越界" in ex.dropped[0]["reason"]),
          ex.dropped[0]["reason"] if ex.dropped else "")

    # 正常值不该被误删
    ex.dropped = []
    obj2 = {"records": [{"row_label": "A1", "fields": {
        "composition.Si": {"raw": "0.62", "quote": "0.62"},
        "composition.Mg": {"raw": "0.85", "quote": "0.85"},
    }}]}
    recs2 = ex._to_records(obj2, [ch], specs)
    check("正常成分保留（不误删）", len(recs2) == 1 and len(ex.dropped) == 0,
          f"留下 {len(recs2)} 条 / 丢弃 {len(ex.dropped)} 条")


# ---------------------------------------------------------------- 20. 结果落盘

def t_result_persist() -> None:
    """关页面 -> 服务自停 -> 重开：上一批结果必须还在（不能只剩「继续抽取」按钮）。"""
    print("\n[20] 结果落盘与重开恢复")
    import tempfile as _tf

    from rag.web import app as webapp

    tmpdir = _tf.mkdtemp(prefix="ragtest_result_")
    p = os.path.join(tmpdir, "last_extract.json")
    orig_path, orig_jobs = webapp._last_result_path, dict(webapp._JOBS)
    orig_done = webapp._RESTORED_DONE
    try:
        webapp._last_result_path = lambda: p
        webapp._JOBS.clear()

        def job(skip, ids):
            return {"started_at": 1000.0, "finished_at": 1100.0,
                    "done_ids": list(ids), "req": {"skip_ids": skip}}

        # 第一批（全新开始）
        webapp._save_last_result("j1", job([], ["c1", "c2"]),
                                 {"headers": ["合金", "Si"], "rows": [{"合金": "A1"}],
                                  "row_ids": ["r1"], "dropped": [], "records": [{}],
                                  "evidence": {"r1": {}}, "records_count": 1,
                                  "picked_chunks": 2})
        disk = webapp._read_last_result()
        check("第一批结果写盘", disk.get("result", {}).get("records_count") == 1
              and disk.get("done_ids") == ["c1", "c2"])

        # 第二批（继续抽取：skip_ids 非空）-> 接在上一批后面
        webapp._save_last_result("j2", job(["c1", "c2"], ["c3"]),
                                 {"headers": ["合金", "Si"], "rows": [{"合金": "A2"}],
                                  "row_ids": ["r2"], "dropped": [], "records": [{}],
                                  "evidence": {"r2": {}}, "records_count": 1,
                                  "picked_chunks": 1})
        disk = webapp._read_last_result()
        check("续抽结果拼接（行数累加，不是覆盖）",
              len(disk["result"]["rows"]) == 2
              and disk["result"]["records_count"] == 2
              and disk["result"]["picked_chunks"] == 3)
        check("已抽块 id 累积去重", disk["done_ids"] == ["c1", "c2", "c3"])

        # 全新开始 -> 整份覆盖
        webapp._save_last_result("j3", job([], ["c9"]),
                                 {"headers": ["合金"], "rows": [{"合金": "B1"}],
                                  "row_ids": ["r9"], "dropped": [], "records": [{}],
                                  "evidence": {}, "records_count": 1,
                                  "picked_chunks": 1})
        check("全新开始 -> 覆盖上一批",
              webapp._read_last_result()["result"]["picked_chunks"] == 1)

        # 模拟进程重启：内存清空，只有盘上这份
        webapp._JOBS.clear()
        webapp._RESTORED_DONE = False
        got = webapp._running_job(now=1200.0)
        check("重开后能找回上一批结果",
              bool(got) and got.get("done") and got.get("result") is not None
              and got.get("result_expired") is False)
        # 捞回来之后就登记成普通任务：再刷新按 TTL 判断，不会每次都重读盘搬运
        later = webapp._running_job(now=1200.0 + webapp._ADOPT_TTL_S + 60)
        check("恢复只做一次（超时后不再搬运，只留已抽块 id）",
              later.get("result") is None and later.get("result_expired") is True
              and bool(later.get("done_ids")))
        check("恢复后登记成已完成任务（之后可被新任务顶掉）",
              any(j.get("done") and j.get("result") for j in webapp._JOBS.values()))

        # 隔太久（超过 web.restore_ttl_s）就不还原：那时多半已经换了一批文献
        webapp._JOBS.clear()
        webapp._RESTORED_DONE = False
        old = webapp._read_last_result()
        old["finished_at"] = 1200.0 - webapp._RESTORE_TTL_S - 1
        with open(p, "w", encoding="utf-8") as f:
            json.dump(old, f, ensure_ascii=False)
        check("超过还原时限 -> 不再把上周的结果摆出来",
              webapp._running_job(now=1200.0) == {})

        webapp._clear_last_result()
        check("清盘后读回空", webapp._read_last_result() == {})
    finally:
        webapp._last_result_path = orig_path
        webapp._JOBS.clear()
        webapp._JOBS.update(orig_jobs)
        webapp._RESTORED_DONE = orig_done
        shutil.rmtree(tmpdir, ignore_errors=True)


# ------------------------------------------------- 21. 写表：列名对不上必须报错不能报成功

def _verified_record():
    from rag.models import Evidence, FieldValue, MaterialRecord
    rec = MaterialRecord(alloy="6013", row_label="6013", source_doc="demo.pdf",
                         source_page=2, record_id="r1")
    for el, v in (("Si", 0.62), ("Mg", 0.85)):
        rec.composition[el] = FieldValue(
            value=v, unit="wt.%", parsed={"kind": "single", "value": v},
            evidence=Evidence(quote=str(v), chunk_id="c1", page=2,
                              doc_name="demo.pdf", verified=True))
    return rec


def t_excel_write_mapping() -> None:
    """导出写表：列名写法不同要能对齐，一列都对不上必须报错而不是假装成功。"""
    print("\n[21] 写表的列名对齐与失败可见性")
    import openpyxl

    from rag.extract.schema import strict_spec_for_header
    from rag.web.app import ExportReq, _map_rows_to_headers, export_data

    key = lambda h: (lambda sp: sp.key if sp else None)(strict_spec_for_header(h))
    check("Si (wt.%) -> composition.Si", key("Si (wt.%)") == "composition.Si")
    check("Mg 含量 -> composition.Mg", key("Mg 含量") == "composition.Mg")
    check("硅 -> composition.Si", key("硅") == "composition.Si")
    check("牌号 -> alloy", key("牌号") == "alloy")
    check("w(Si) -> composition.Si", key("w(Si)") == "composition.Si")
    check("'其余' 这种含糊表头不硬猜 -> None", key("其余") is None)

    rec = _verified_record()
    sd_yaml = open(os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                                "data", "schemas", "al_no_balance.yaml"),
                   encoding="utf-8").read()
    tmpdir = tempfile.mkdtemp(prefix="ragtest_write_")

    def export(headers, name, make_file=True):
        p = os.path.join(tmpdir, f"{name}.xlsx")
        if make_file:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "data"
            if headers:
                ws.append(headers)
            wb.save(p)
            wb.close()
        res = export_data(ExportReq(records=[rec.to_dict()], schema_text=sd_yaml,
                                    format="xlsx", excel_path=p, clean={}))
        body = res if isinstance(res, dict) else json.loads(res.body.decode())
        grid = [tuple(r) for r in openpyxl.load_workbook(p).active.iter_rows(values_only=True)]
        return body, grid

    # 关键回归：表头带单位时原来一个格子都写不进去，却照样返回"成功"
    body, grid = export(["合金牌号", "Si (wt.%)", "Mg (wt.%)"], "unit")
    check("表头带单位 -> 真的写进去了",
          int(body.get("written_cells") or 0) > 0
          and grid[1][1] == 0.62 and grid[1][2] == 0.85, str(grid[:2]))
    check("改名对齐要在返回值里可查",
          any("Si -> Si (wt.%)" in m for m in body.get("column_mapping") or []))

    body, grid = export(["牌号", "硅", "镁"], "alias")
    check("中文别名表头 -> 写进对应列",
          int(body.get("written_cells") or 0) == 3
          and grid == [("牌号", "硅", "镁"), ("6013", 0.62, 0.85)], str(grid))

    body, grid = export(["序号", "备注", "日期"], "unrelated")
    check("一列都对不上 -> 明确失败（不是假成功）",
          body.get("ok") is False and body.get("error"), str(body.get("ok")))
    check("失败时不许改动用户的表", grid == [("序号", "备注", "日期")], str(grid))

    body, grid = export([], "empty")
    check("空表（无表头）-> 自动建表头再写数据",
          grid[0][:3] == ("合金牌号", "Si", "Mg"), str(grid[0]))

    # 纯函数层：模板列 -> 目标列名
    specs = {"alloy": None}
    from rag.extract.schema import BY_KEY
    sps = {"alloy": BY_KEY["alloy"], "composition.Si": BY_KEY["composition.Si"]}
    mapped, mapping, skipped = _map_rows_to_headers(
        [{"合金牌号": "A", "Si": 0.6}], sps, ["牌号", "硅"])
    check("行改名：模板列换为目标表真实列名",
          mapped == [{"牌号": "A", "硅": 0.6}] and len(mapping) == 2 and skipped == [],
          f"{mapped} / {mapping} / {skipped}")

    shutil.rmtree(tmpdir, ignore_errors=True)


def main() -> int:
    tmp = tempfile.mkdtemp(prefix="ragtest_")
    try:
        t_parse_value()
        t_verify()
        t_schema()
        t_two_column(tmp)
        t_pipeline(tmp)
        t_extract_verify()
        t_merge_records()
        t_table_context()
        t_excel(tmp)
        t_format_text()
        t_rules()
        t_schemas(tmp)
        t_clean()
        t_chunk_plan(tmp)
        t_launcher()
        t_job_stream()
        t_adopt_ttl()
        t_alignment()
        t_sanity_range()
        t_result_persist()
        t_excel_write_mapping()
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

    print(f"\n===== 汇总：通过 {len(PASS)} 项，失败 {len(FAIL)} 项 =====")
    if FAIL:
        for f in FAIL:
            print(f"  失败：{f}")
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
